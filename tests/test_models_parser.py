from io import BytesIO
import json
from pathlib import Path
import zipfile

import pytest
from pydantic import ValidationError

from ahb.models import Entry, Link, Profile, Site, blank_site
from ahb.parser import ResumeError, parse_resume, parse_text


def pdf_bytes(*, text=None, encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=100)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        content = DecodedStreamObject()
        content.set_data(f'BT /F1 12 Tf 20 60 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = content
    if encrypted:
        writer.encrypt('private-test-password')
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_blank_and_bilingual_demo_validate_without_personal_defaults():
    blank = blank_site()
    assert blank.profile.email == '' and blank.profile.links == []
    assert not blank.example
    assert not any(blank.sections.values())
    demo = Site.model_validate_json((Path(__file__).parents[1] / 'examples/demo.json').read_text(encoding='utf8'))
    assert demo.example and set(demo.languages) == {'en', 'zh'}
    assert demo.site_url == ''
    assert demo.profile.email.endswith('@example.org')


@pytest.mark.parametrize('url', ['javascript:alert(1)', 'data:text/html,hello', '//example.org', 'http://example.org', 'https://user:password@example.org', 'https://example.org\\@attacker.example', 'https://example.org/\nX-Test:yes', 'https://example.org:99999'])
def test_dangerous_urls_rejected_before_preview(url):
    with pytest.raises(ValidationError):
        Link(label='Unsafe', url=url)
    with pytest.raises(ValidationError):
        Entry(url=url)


@pytest.mark.parametrize('path', ['../private.png', 'assets/uploads/../secret.png', 'assets/uploads/%2e%2e.png', '/assets/uploads/photo.png', 'https://example.org/photo.png', 'assets/uploads/photo.svg', 'assets/uploads/photo.png?x=1', 'assets/uploads/document.pdf'])
def test_image_paths_are_uploaded_raster_only(path):
    with pytest.raises(ValidationError):
        Profile(portrait=path)


def test_public_cv_explicit_pdf_and_no_unknown_private_fields():
    assert Profile(cv='assets/uploads/approved-cv.pdf').cv.endswith('.pdf')
    with pytest.raises(ValidationError):
        Profile(cv='assets/uploads/avatar.png')
    with pytest.raises(ValidationError):
        Profile(home_address='Not a public field')
    with pytest.raises(ValidationError):
        Profile(email='public@example.org\r\nBcc:someone@example.org')
    with pytest.raises(ValidationError):
        Site(googleSiteVerification='Unrecognized carry-over data')


def test_languages_origin_and_section_shape():
    assert Site(languages=['zh'], default_language='zh').languages == ['zh']
    assert Site(site_url='https://sample.pages.dev').site_url == 'https://sample.pages.dev/'
    for bad in ['https://example.org/subdir', 'https://example.org/?q=1', 'https://example.org/#home']:
        with pytest.raises(ValidationError):
            Site(site_url=bad)
    with pytest.raises(ValidationError):
        Site(languages=['zh'])
    with pytest.raises(ValidationError):
        Site(languages=['en', 'en'])
    with pytest.raises(ValidationError):
        Site(sections={'education': [{}] * 101})
    with pytest.raises(ValidationError):
        Profile(name={'en': 'a' * 1001})


def test_parser_candidates_remain_unconfirmed_and_does_not_invent_authorship():
    result = parse_text('Alex Morgan\nalex@example.org\nhttps://github.com/example\nEducation\nExample University\n2023.09 – 2027.06\n\nPublications\nAn Example Title\nUnder review\nAuthors: Alex Morgan, Taylor Chen\nProjects\nA Demo Project\nA local research notebook.')
    site = Site.model_validate(result['site'])
    assert site.profile.name['en'] == 'Alex Morgan'
    assert site.profile.email == 'alex@example.org'
    assert site.profile.links[0].label == 'GitHub'
    assert site.sections['education'][0].date == '2023.09 – 2027.06'
    paper = site.sections['publications'][0]
    assert not paper.confirmed and paper.authors == '' and not any(paper.status.values())
    assert 'Under review' in paper.description['en']
    assert result['warnings']


def test_chinese_resume_and_no_unrequested_translation():
    result = parse_text('林知远\n个人简介\n我关注科研工具的可访问性。\n教育背景\n示例大学本科\n2021 – 2025\n项目经历\n开放研究笔记\n用于整理研究资料。')
    site = Site.model_validate(result['site'])
    assert site.default_language == 'zh'
    assert site.profile.name['zh'] == '林知远'
    assert site.profile.name['en'] == ''
    assert site.sections['education'][0].title == {'zh': '示例大学本科'}


def test_phone_address_and_id_are_not_promoted_to_profile():
    result = parse_text('Alex Morgan\nPhone: +1 555 0100\nAddress: 1 Example Street\nNational ID: DEMO-NOT-REAL\nSkills\nPython')
    public = json.dumps(result['site'])
    assert '555' not in public and 'Example Street' not in public and 'DEMO-NOT-REAL' not in public
    assert 'Phone:' in result['text']  # raw text stays separate for local review


def test_pdf_text_and_unreadable_scan_encrypted_corrupt_paths():
    assert 'Alex Morgan' in parse_resume('resume.pdf', pdf_bytes(text='Alex Morgan'))['text']
    with pytest.raises(ResumeError, match='scanned'):
        parse_resume('scan.pdf', pdf_bytes())
    with pytest.raises(ResumeError, match='encrypted'):
        parse_resume('locked.pdf', pdf_bytes(encrypted=True))
    with pytest.raises(ResumeError, match='not a valid PDF'):
        parse_resume('fake.pdf', b'not a pdf')
    with pytest.raises(ResumeError, match='could not be read'):
        parse_resume('broken.pdf', b'%PDF-1.7\ncorrupt')


def test_docx_extracts_paragraphs_and_table_text_in_document_order():
    from docx import Document
    document = Document()
    document.add_paragraph('Alex Morgan')
    document.add_paragraph('Education')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Example University'
    table.cell(0, 1).text = '2023 – 2027'
    document.add_paragraph('Projects')
    document.add_paragraph('A fictional study notebook')
    buffer = BytesIO()
    document.save(buffer)
    result = parse_resume('resume.docx', buffer.getvalue())
    assert result['text'].index('Education') < result['text'].index('Example University') < result['text'].index('Projects')
    assert result['site']['sections']['education'][0]['confirmed'] is False
    with pytest.raises(ResumeError):
        parse_resume('bad.docx', b'not a zip')


def test_docx_rejects_entity_declarations_before_loading_document():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, 'w') as z:
        z.writestr('[Content_Types].xml', '<Types/>')
        z.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY y SYSTEM "file:///private">]><x/>')
    with pytest.raises(ResumeError, match='XML declarations'):
        parse_resume('entity.docx', buffer.getvalue())


def test_file_and_text_resource_limits():
    with pytest.raises(ResumeError, match='20 MB'):
        parse_resume('large.txt', b'A' * (20 * 1024 * 1024 + 1))
    with pytest.raises(ResumeError, match='120,000'):
        parse_text('A' * 120001)
    with pytest.raises(ResumeError, match='Supported'):
        parse_resume('resume.exe', b'hello')
    with pytest.raises(ResumeError, match='UTF-8'):
        parse_resume('bad.txt', b'\xff\xff\xff')
    assert parse_resume('unicode.txt', '林知远'.encode('utf-16'))['text'] == '林知远'


def test_docx_archive_bomb_and_path_are_rejected():
    for name, data in [('word/payload.bin', b'0' * (6 * 1024 * 1024)), ('../private.xml', b'<x/>')]:
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr('[Content_Types].xml', '<Types/>')
            z.writestr('word/document.xml', '<x/>')
            z.writestr(name, data)
        with pytest.raises(ResumeError, match='unsafe'):
            parse_resume('unsafe.docx', buffer.getvalue())
